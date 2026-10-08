import json
import os
from datetime import date

import pytest
from openpyxl import load_workbook

from exportar import (
    ExportacionCancelada, convertir_celda, exportar, exportar_incremental, nombre_hoja,
)

COLS = ["empresa", "valor", "fecha", "url", "objeto"]
ENC = ["Empresa", "Valor", "Fecha", "URL", "Objeto"]
TIT = dict(zip(COLS, ENC))
FILAS = [
    {"empresa": "ACME", "valor": "$1.500.000", "fecha": "2024-05-01",
     "url": "https://x.test/a", "objeto": "a" * 40000},
    {"empresa": "Beta", "valor": "", "fecha": "", "url": "", "objeto": "mal\x00caracter"},
]


def test_xlsx_con_formato(tmp_path):
    ruta = tmp_path / "a.xlsx"
    exportar("xlsx", COLS, FILAS, str(ruta), TIT)
    ws = load_workbook(ruta)["Resumen"]
    assert ws["A1"].value == "Empresa" and ws["A1"].font.bold
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref
    assert ws["B2"].value == 1500000 and "$" in ws["B2"].number_format
    assert ws["C2"].value.date() == date(2024, 5, 1)
    assert ws["D2"].hyperlink.target == "https://x.test/a"
    assert len(ws["E2"].value) == 32767          # limite de Excel
    assert ws["E3"].value == "malcaracter"        # caracter de control eliminado
    assert ws["B3"].value is None and ws["D3"].hyperlink is None
    assert ws.column_dimensions["D"].width >= 40     # url (D y E se agrupan si tienen el mismo ancho)


def test_xlsx_numero_en_texto_de_columna_valor(tmp_path):
    ruta = tmp_path / "b.xlsx"
    exportar("xlsx", ["precio_base"], [{"precio_base": "65000000"}], str(ruta))
    assert load_workbook(ruta)["Resumen"]["A2"].value == 65000000


def test_convertir_celda():
    assert convertir_celda("valor", "$2.000") == (2000, '"$"#,##0')
    assert convertir_celda("x", "2024-05-01")[0] == date(2024, 5, 1)
    assert convertir_celda("x", "mal\x00") == ("mal", None)
    assert convertir_celda("x", 7) == (7, None)
    assert convertir_celda("x", None) == (None, None)
    assert convertir_celda("x", {"url": "https://u"}) == ("https://u", None)
    # Las hojas de detalle traen marcas ISO de Socrata: deben quedar como fecha real
    assert convertir_celda("fecha_de_firma", "2025-03-01T00:00:00.000") == (
        date(2025, 3, 1), "yyyy-mm-dd")


def test_nombres_de_hoja_validos_y_unicos():
    usados = set()
    a = nombre_hoja("x" * 40, usados)
    b = nombre_hoja("x" * 40, usados)
    assert a != b and len(a) <= 31 and len(b) <= 31
    assert not set("[]:*?/\\") & set(nombre_hoja("a/b:c[d]", set()))


def test_csv_utf8_bom_y_delimitador(tmp_path):
    ruta = tmp_path / "a.csv"
    exportar("csv", ["empresa", "valor"], FILAS[:1], str(ruta),
             {"empresa": "Empresa", "valor": "Valor"}, delimitador=";")
    crudo = open(ruta, "rb").read()
    assert crudo.startswith(b"\xef\xbb\xbf")
    assert crudo.decode("utf-8-sig").splitlines() == ["Empresa;Valor", "ACME;1500000"]


def test_csv_celdas_con_saltos_y_comillas_ida_y_vuelta(tmp_path):
    import csv
    ruta = tmp_path / "c.csv"
    exportar("csv", ["a"], [{"a": 'linea1\nlinea2 "x"'}, {"a": "ñandú"}], str(ruta))
    with open(ruta, encoding="utf-8-sig", newline="") as f:
        assert list(csv.reader(f)) == [["a"], ['linea1\nlinea2 "x"'], ["ñandú"]]


def test_json(tmp_path):
    ruta = tmp_path / "a.json"
    exportar("json", ["empresa", "valor"], FILAS[:1], str(ruta))
    assert json.load(open(ruta, encoding="utf-8")) == [{"empresa": "ACME", "valor": 1500000}]


def test_csv_por_dataset(tmp_path):
    rutas = exportar("csv_dataset", ["a"], [{"a": "1"}], str(tmp_path), None,
                     {"Contratos SECOP II": [{"b": "2"}]}, sello="20260101_1200")
    assert len(rutas) == 2 and all(os.path.exists(r) for r in rutas)
    assert "Contratos_SECOP_II" in os.path.basename(rutas[1])


def test_xlsx_agrega_hojas_de_detalle_con_union_de_columnas(tmp_path):
    ruta = tmp_path / "t.xlsx"
    detalle = {"Procesos SECOP II": [{"id": "1", "urlproceso": {"url": "https://p.test/1"}},
                                      {"id": "2", "extra": "z"}]}
    exportar("xlsx", ["empresa"], [{"empresa": "A"}], str(ruta), {"empresa": "Empresa"}, detalle)
    wb = load_workbook(ruta)
    assert wb.sheetnames == ["Resumen", "Procesos SECOP II"]
    ws = wb["Procesos SECOP II"]
    assert [c.value for c in ws[1]] == ["id", "urlproceso", "extra"]
    assert ws["B2"].value == "https://p.test/1" and ws["B2"].hyperlink.target == "https://p.test/1"


def test_sin_columnas_o_sin_filas_no_crea_archivo(tmp_path):
    ruta = tmp_path / "x.csv"
    with pytest.raises(ValueError):
        exportar("csv", [], [{"a": 1}], str(ruta))
    with pytest.raises(ValueError):
        exportar("csv", ["a"], [], str(ruta))
    assert not ruta.exists()


def test_formato_desconocido(tmp_path):
    with pytest.raises(ValueError):
        exportar("pdf", ["a"], [{"a": 1}], str(tmp_path / "x"))


# --- incremental ---------------------------------------------------------
def test_incremental_acumula_paginas_sin_perder_columnas_de_detalle(tmp_path):
    def paginas():
        yield [{"a": "1"}], {"D": [{"x": 1}]}
        yield [{"a": "2"}], {"D": [{"y": 2}]}
    ruta = tmp_path / "i.xlsx"
    archivos, n = exportar_incremental("xlsx", ["a"], str(ruta), paginas())
    assert archivos == [str(ruta)] and n == 2
    wb = load_workbook(ruta)
    assert [c.value for c in wb["Resumen"]["A"]] == ["a", "1", "2"]
    assert [c.value for c in wb["D"][1]] == ["x", "y"]


def test_incremental_acepta_un_generador_y_cuenta_las_filas(tmp_path):
    consumidas = []

    def paginas():
        for i in range(3):
            consumidas.append(i)
            yield [{"a": str(i)}], {}
    pendientes = paginas()
    archivos, n = exportar_incremental("csv", ["a"], str(tmp_path / "p.csv"), pendientes)
    assert consumidas == [0, 1, 2] and n == 3


def test_cancelar_durante_la_descarga_no_crea_nada(tmp_path):
    ruta = tmp_path / "c.csv"

    def paginas():
        yield [{"a": "1"}], {}
        yield [{"a": "2"}], {}
    estado = {"n": 0}

    def cancelado():
        estado["n"] += 1
        return estado["n"] >= 2
    with pytest.raises(ExportacionCancelada):
        exportar_incremental("csv", ["a"], str(ruta), paginas(), cancelado=cancelado)
    assert not ruta.exists()


def test_cancelar_durante_la_escritura_borra_el_archivo_parcial(tmp_path):
    ruta = tmp_path / "c.csv"
    filas = [{"a": str(i)} for i in range(2500)]
    llamadas = {"n": 0}

    def cancelado():
        llamadas["n"] += 1
        return llamadas["n"] > 3      # deja pasar la descarga y cancela escribiendo (fila 1000)
    with pytest.raises(ExportacionCancelada):
        exportar_incremental("csv", ["a"], str(ruta), iter([(filas, {})]), cancelado=cancelado)
    assert not ruta.exists()


def test_error_al_descargar_no_deja_archivo(tmp_path):
    def paginas():
        yield [{"a": "1"}], {}
        raise RuntimeError("red caida")
    ruta = tmp_path / "e.csv"
    with pytest.raises(RuntimeError):
        exportar_incremental("csv", ["a"], str(ruta), paginas())
    assert not ruta.exists()


def test_error_no_borra_un_archivo_preexistente_si_no_llego_a_escribir(tmp_path):
    ruta = tmp_path / "previo.csv"
    ruta.write_text("conservar", encoding="utf-8")

    def paginas():
        raise RuntimeError("red caida")
        yield
    with pytest.raises(RuntimeError):
        exportar_incremental("csv", ["a"], str(ruta), paginas())
    assert ruta.read_text(encoding="utf-8") == "conservar"


# --- no tocar archivos preexistentes ---------------------------------------
def _cancela_escribiendo():
    llamadas = {"n": 0}

    def cancelado():
        llamadas["n"] += 1
        return llamadas["n"] > 3
    return cancelado


@pytest.mark.parametrize("formato,ext", [("xlsx", "xlsx"), ("csv", "csv"), ("json", "json")])
def test_cancelar_escribiendo_conserva_archivo_previo(tmp_path, formato, ext):
    ruta = tmp_path / f"previo.{ext}"
    ruta.write_bytes(b"original")
    filas = [{"a": str(i)} for i in range(2500)]
    with pytest.raises(ExportacionCancelada):
        exportar_incremental(formato, ["a"], str(ruta), iter([(filas, {})]),
                             cancelado=_cancela_escribiendo())
    assert ruta.read_bytes() == b"original"
    assert [p.name for p in tmp_path.iterdir()] == [ruta.name]


@pytest.mark.parametrize("formato,ext", [("xlsx", "xlsx"), ("csv", "csv"), ("json", "json")])
def test_exito_reemplaza_archivo_previo_sin_dejar_part(tmp_path, formato, ext):
    ruta = tmp_path / f"previo.{ext}"
    ruta.write_bytes(b"original")
    archivos, _ = exportar_incremental(formato, ["a"], str(ruta), iter([([{"a": "1"}], {})]))
    assert archivos == [str(ruta)] and ruta.read_bytes() != b"original"
    assert [p.name for p in tmp_path.iterdir()] == [ruta.name]


def test_csv_dataset_no_sobrescribe_existentes(tmp_path):
    previo = tmp_path / "SECOP_Resumen_s.csv"
    previo.write_text("conservar", encoding="utf-8")
    rutas = exportar("csv_dataset", ["a"], [{"a": "1"}], str(tmp_path), sello="s")
    assert previo.read_text(encoding="utf-8") == "conservar"
    assert len(rutas) == 1 and rutas[0] != str(previo) and os.path.exists(rutas[0])


def test_csv_dataset_cancelar_limpia_lo_creado(tmp_path):
    filas = [{"a": str(i)} for i in range(2500)]
    detalle = {"D": [{"b": str(i)} for i in range(2500)]}
    llamadas = {"n": 0}

    def cancelado():
        llamadas["n"] += 1
        return llamadas["n"] > 5   # resumen se escribe completo; cancela en el detalle
    with pytest.raises(ExportacionCancelada):
        exportar_incremental("csv_dataset", ["a"], str(tmp_path), iter([(filas, detalle)]),
                             sello="s", cancelado=cancelado)
    assert list(tmp_path.iterdir()) == []


def test_cierra_el_generador_de_paginas(tmp_path):
    estado = {"cerrado": False}

    def paginas():
        try:
            yield [{"a": "1"}], {}
            yield [{"a": "2"}], {}
        finally:
            estado["cerrado"] = True
    with pytest.raises(ExportacionCancelada):
        exportar_incremental("csv", ["a"], str(tmp_path / "g.csv"), paginas(),
                             cancelado=lambda: True)
    assert estado["cerrado"]
