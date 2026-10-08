"""Pestanas de detalle (Contratos, Procesos, Integrado): orden, abrir en SECOP y exportar."""
import csv
import json
import tkinter as tk
from tkinter import ttk

import pytest
from openpyxl import load_workbook

import tabla_detalle
from tabla_detalle import TablaDetalle

URL_A = "https://community.secop.gov.co/Public/Tendering/OpportunityDetail/Index?noticeUID=A"
URL_B = "https://community.secop.gov.co/Public/Tendering/OpportunityDetail/Index?noticeUID=B"

CONTRATOS = [
    {"id_contrato": "C1", "valor_del_contrato": "5000", "nombre_entidad": "beta",
     "fecha_de_firma": "2024-03-01T00:00:00.000", "urlproceso": {"url": URL_A}},
    {"id_contrato": "C2", "valor_del_contrato": "100000", "nombre_entidad": "Alfa",
     "fecha_de_firma": "2023-12-31T00:00:00.000", "urlproceso": {"url": URL_B}},
    {"id_contrato": "C3", "nombre_entidad": ""},            # Socrata omite los nulos
]
PROCESOS = [
    {"id_del_proceso": "P1", "urlproceso": {"url": URL_A}},
    {"id_del_proceso": "P2", "urlproceso": {"url": "javascript:alert(1)"}},
    {"id_del_proceso": "P3"},
]
INTEGRADO = [
    {"numero_del_contrato": "I1", "url_contrato": URL_B},
    {"numero_del_contrato": "I2", "url_contrato": "file:///C:/Windows/system32/calc.exe"},
]


def _tabla(app, nombre, data):
    app._reconstruir_tabs_detalle({nombre: data})
    return app._tablas_detalle[nombre]


def _col(tabla, col):
    return [tabla.tree.set(i, col) for i in tabla.tree.get_children()]


def _item(tabla, col, valor):
    return next(i for i in tabla.tree.get_children() if tabla.tree.set(i, col) == valor)


@pytest.fixture
def abiertas(monkeypatch):
    urls = []
    monkeypatch.setattr(tabla_detalle.webbrowser, "open", lambda url, new=0: urls.append(url))
    return urls


# ---- orden -------------------------------------------------------------------

def test_cada_pestana_de_detalle_tiene_su_tabla(app):
    app._reconstruir_tabs_detalle({"Contratos SECOP II": CONTRATOS, "Procesos SECOP II": PROCESOS,
                                   "SECOP Integrado (historico)": INTEGRADO})
    assert set(app._tablas_detalle) == {"Contratos SECOP II", "Procesos SECOP II",
                                        "SECOP Integrado (historico)"}
    app._destruir_tabs_detalle()
    assert app._tablas_detalle == {}


def test_ordenar_numeros_por_valor_real_y_vacios_al_final(app):
    t = _tabla(app, "Contratos SECOP II", CONTRATOS)
    t.ordenar("valor_del_contrato")
    assert _col(t, "id_contrato") == ["C1", "C2", "C3"]     # 5.000 < 100.000 (no como texto)
    t.ordenar("valor_del_contrato")
    assert _col(t, "id_contrato") == ["C2", "C1", "C3"]     # invertido; el vacio sigue al final


def test_ordenar_texto_sin_distinguir_mayusculas_y_fechas(app):
    t = _tabla(app, "Contratos SECOP II", CONTRATOS)
    t.ordenar("nombre_entidad")
    assert _col(t, "nombre_entidad") == ["Alfa", "beta", ""]
    t.ordenar("nombre_entidad")
    assert _col(t, "nombre_entidad") == ["beta", "Alfa", ""]
    t.ordenar("fecha_de_firma")
    assert _col(t, "id_contrato") == ["C2", "C1", "C3"]


def test_ordenar_conserva_cebra_y_filas_completas_para_copiar(app):
    largo = "x" * 300
    data = [{"objeto": largo, "n": "2"}, {"objeto": "corto", "n": "1"}]
    t = _tabla(app, "Procesos SECOP II", data)
    t.ordenar("n")
    items = t.tree.get_children()
    assert [t.tree.item(i, "tags")[0] for i in items] == ["par", "impar"]
    assert len(t.tree.set(items[1], "objeto")) <= 120         # pantalla recortada
    assert t.copiador.filas_visibles()[1]["objeto"] == largo  # copia completa, en el orden nuevo


def test_el_encabezado_ordena_al_hacer_clic(app):
    t = _tabla(app, "Contratos SECOP II", CONTRATOS)
    assert t.tree.heading("valor_del_contrato", "command")


# ---- abrir en SECOP ------------------------------------------------------------

def test_url_de_contratos_procesos_e_integrado(app):
    c = _tabla(app, "Contratos SECOP II", CONTRATOS)
    assert c.url_de_item(_item(c, "id_contrato", "C1")) == URL_A
    assert c.url_de_item(_item(c, "id_contrato", "C3")) == ""
    p = _tabla(app, "Procesos SECOP II", PROCESOS)
    assert p.url_de_item(_item(p, "id_del_proceso", "P1")) == URL_A     # dict {"url": ...}
    assert p.url_de_item(_item(p, "id_del_proceso", "P2")) == ""        # no http(s)
    i = _tabla(app, "SECOP Integrado (historico)", INTEGRADO)
    assert i.url_de_item(_item(i, "numero_del_contrato", "I1")) == URL_B
    assert i.url_de_item(_item(i, "numero_del_contrato", "I2")) == ""


def test_doble_clic_abre_la_fila_seleccionada(app, abiertas):
    t = _tabla(app, "SECOP Integrado (historico)", INTEGRADO)
    t.tree.selection_set(_item(t, "numero_del_contrato", "I1"))
    t._on_doble_clic(None)
    assert abiertas == [URL_B]
    assert URL_B in app.lbl_estado.cget("text")


def test_no_se_abren_urls_que_no_son_http(app, abiertas):
    t = _tabla(app, "Procesos SECOP II", PROCESOS)
    t.tree.selection_set(_item(t, "id_del_proceso", "P2"))
    t._on_doble_clic(None)
    t.abrir_en_secop(_item(t, "id_del_proceso", "P3"))
    assert abiertas == []
    assert "no tiene URL" in app.lbl_estado.cget("text")
    assert tabla_detalle.abrir_url("javascript:alert(1)") is False
    assert abiertas == []


def test_el_orden_mantiene_la_url_de_cada_fila(app, abiertas):
    t = _tabla(app, "Contratos SECOP II", CONTRATOS)
    t.ordenar("valor_del_contrato")
    t.ordenar("valor_del_contrato")                         # C2 primero
    assert t.url_de_item(t.tree.get_children()[0]) == URL_B


def _etiquetas(menu):
    fin = menu.index("end")
    return [menu.entrycget(i, "label") if menu.type(i) == "command" else "-"
            for i in range(fin + 1)] if fin is not None else []


def test_menu_contextual_abre_en_secop_y_exporta(app, abiertas):
    t = _tabla(app, "Contratos SECOP II", CONTRATOS)
    menu = tk.Menu(t.tree, tearoff=0)
    t._menu_extra(menu, _item(t, "id_contrato", "C1"))
    etiquetas = _etiquetas(menu)
    assert "Abrir en SECOP" in etiquetas and "Exportar esta tabla..." in etiquetas
    menu.invoke(etiquetas.index("Abrir en SECOP"))
    assert abiertas == [URL_A]
    sin_url = tk.Menu(t.tree, tearoff=0)
    t._menu_extra(sin_url, _item(t, "id_contrato", "C3"))
    assert str(sin_url.entrycget(_etiquetas(sin_url).index("Abrir en SECOP"), "state")) == "disabled"


def test_el_copiador_usa_el_menu_extra(app):
    t = _tabla(app, "Contratos SECOP II", CONTRATOS)
    assert t.copiador._extra_menu == t._menu_extra


def test_cursor_de_mano_solo_sobre_filas_con_url(app, monkeypatch):
    t = _tabla(app, "Contratos SECOP II", CONTRATOS)
    con, sin = _item(t, "id_contrato", "C1"), _item(t, "id_contrato", "C3")
    evento = type("E", (), {"x": 1, "y": 1})()
    monkeypatch.setattr(t.tree, "identify_row", lambda y: con)
    t._cursor_hover(evento)
    assert str(t.tree.cget("cursor")) == "hand2" and URL_A in app.lbl_estado.cget("text")
    monkeypatch.setattr(t.tree, "identify_row", lambda y: sin)
    t._cursor_hover(evento)
    assert str(t.tree.cget("cursor")) == ""


def test_texto_de_ayuda_al_pie(app):
    t = _tabla(app, "Contratos SECOP II", CONTRATOS)
    texto = t.lbl_ayuda.cget("text")
    assert "3 registros" in texto
    for parte in ("Doble clic: abrir en SECOP", "Ctrl+C", "Clic derecho", "Clic en columna: ordenar"):
        assert parte in texto


# ---- exportar ------------------------------------------------------------------

@pytest.fixture
def dialogos(monkeypatch):
    avisos = {"info": [], "error": []}
    monkeypatch.setattr(tabla_detalle.messagebox, "showinfo",
                        lambda *a, **k: avisos["info"].append(a))
    monkeypatch.setattr(tabla_detalle.messagebox, "showerror",
                        lambda *a, **k: avisos["error"].append(a))
    return avisos


@pytest.fixture
def salida(tmp_path):
    """Carpeta propia: tmp_path tambien es el cwd y la carpeta de datos de la app."""
    carpeta = tmp_path / "salida"
    carpeta.mkdir()
    return carpeta


def _destino(monkeypatch, ruta):
    monkeypatch.setattr(tabla_detalle.filedialog, "asksaveasfilename", lambda **k: str(ruta))


def test_exportar_xlsx_con_todas_las_columnas(app, monkeypatch, salida, dialogos):
    t = _tabla(app, "Contratos SECOP II", CONTRATOS)
    ruta = salida / "contratos.xlsx"
    _destino(monkeypatch, ruta)
    t.exportar_tabla()
    wb = load_workbook(ruta)
    ws = wb[wb.sheetnames[0]]
    assert wb.sheetnames == ["Contratos SECOP II"]
    filas = list(ws.values)
    assert list(filas[0]) == t.titulos and len(t.titulos) == 5
    assert len(filas) == 4
    col_url = t.columnas.index("urlproceso") + 1
    celda = ws.cell(row=2, column=col_url)
    assert celda.value == URL_A and celda.hyperlink.target == URL_A
    assert "contratos.xlsx" in app.lbl_estado.cget("text") and dialogos["error"] == []


def test_exportar_csv_y_json(app, monkeypatch, salida, dialogos):
    t = _tabla(app, "Procesos SECOP II", PROCESOS)
    _destino(monkeypatch, salida / "p.csv")
    t.exportar_tabla()
    with open(salida / "p.csv", encoding="utf-8-sig", newline="") as f:
        filas = list(csv.reader(f))
    assert filas[0] == t.titulos and [r[0] for r in filas[1:]] == ["P1", "P2", "P3"]
    assert filas[1][1] == URL_A                             # el dict se exporta como su URL
    _destino(monkeypatch, salida / "p.json")
    t.exportar_tabla()
    datos = json.loads((salida / "p.json").read_text(encoding="utf-8"))
    assert [d["id_del_proceso"] for d in datos] == ["P1", "P2", "P3"]
    assert set(datos[0]) == set(t.columnas)


def test_exportar_respeta_el_orden_en_pantalla(app, monkeypatch, salida, dialogos):
    t = _tabla(app, "Contratos SECOP II", CONTRATOS)
    t.ordenar("valor_del_contrato")
    t.ordenar("valor_del_contrato")
    _destino(monkeypatch, salida / "c.json")
    t.exportar_tabla()
    datos = json.loads((salida / "c.json").read_text(encoding="utf-8"))
    assert [d["id_contrato"] for d in datos] == ["C2", "C1", "C3"]


def test_exportar_tabla_vacia_avisa_sin_crear_archivo(app, monkeypatch, salida, dialogos):
    monkeypatch.setattr(tabla_detalle.filedialog, "asksaveasfilename",
                        lambda **k: pytest.fail("no debia pedir archivo"))
    t = TablaDetalle(app, ttk.Frame(app), "Vacia", [], lambda c, v: v, lambda m: None)
    t.exportar_tabla()
    assert len(dialogos["info"]) == 1 and list(salida.iterdir()) == []


def test_exportar_cancelado_no_escribe(app, monkeypatch, salida, dialogos):
    t = _tabla(app, "Contratos SECOP II", CONTRATOS)
    _destino(monkeypatch, "")
    t.exportar_tabla()
    assert list(salida.iterdir()) == [] and dialogos == {"info": [], "error": []}


def test_exportar_extension_no_soportada_avisa(app, monkeypatch, salida, dialogos):
    t = _tabla(app, "Contratos SECOP II", CONTRATOS)
    _destino(monkeypatch, salida / "x.txt")
    t.exportar_tabla()
    assert not (salida / "x.txt").exists() and len(dialogos["error"]) == 1


def test_error_al_exportar_se_informa(app, monkeypatch, salida, dialogos):
    def falla(*a, **k):
        raise OSError("disco lleno")
    monkeypatch.setattr(tabla_detalle, "exportar", falla)
    t = _tabla(app, "Contratos SECOP II", CONTRATOS)
    _destino(monkeypatch, salida / "x.xlsx")
    t.exportar_tabla()
    assert len(dialogos["error"]) == 1 and "disco lleno" in dialogos["error"][0][1]
    assert "disco lleno" in app.lbl_estado.cget("text")
    assert list(salida.iterdir()) == []


@pytest.mark.parametrize("ruta,formato", [("a.xlsx", "xlsx"), ("a.CSV", "csv"),
                                          ("a.json", "json"), ("a.txt", None), ("a", None)])
def test_formato_por_extension(ruta, formato):
    assert tabla_detalle.formato_por_extension(ruta) == formato
