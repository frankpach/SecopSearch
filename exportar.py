# -*- coding: utf-8 -*-
"""Exportacion incremental: descarga pagina a pagina a archivos temporales en
disco y escribe el destino leyendo de ahi (xlsx con write_only, csv, json).
Nunca mantiene todas las filas en memoria. Sin Tk ni red."""
import csv
import json
import os
import re
import tempfile

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from tabla_utils import parse_fecha, valor_crudo

_ILEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_NUM_RE = re.compile(r"^-?\d+(\.\d+)?$")
FORMATO_PESOS = '"$"#,##0'
FORMATO_FECHA = "yyyy-mm-dd"
MAX_CELDA = 32767
COLUMNAS_URL = ("urlproceso", "url_contrato", "url")
FORMATOS = ("xlsx", "csv", "json", "csv_dataset")
LOTE_CANCELACION = 1000
NEGRITA = Font(bold=True)
ENLACE = Font(color="0563C1", underline="single")


class ExportacionCancelada(Exception):
    """El usuario cancelo; los archivos parciales ya se borraron."""


def nombre_hoja(nombre, usados):
    base = re.sub(r"[\[\]:*?/\\]", "-", str(nombre)).strip()[:31] or "Hoja"
    cand, i = base, 2
    while cand.lower() in usados:
        suf = f" ({i})"
        cand = base[: 31 - len(suf)] + suf
        i += 1
    usados.add(cand.lower())
    return cand


def convertir_celda(columna, valor):
    """Devuelve (valor para Excel, number_format o None)."""
    if valor is None or valor == "":
        return None, None
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        return valor, None
    crudo = valor_crudo(valor)
    if isinstance(crudo, int):
        return crudo, FORMATO_PESOS
    texto = str(crudo)
    fecha = parse_fecha(texto)
    if fecha is not None:
        return fecha, FORMATO_FECHA
    nombre = columna.lower()
    if ("valor" in nombre or "precio" in nombre) and _NUM_RE.match(texto.strip()):
        return float(texto), FORMATO_PESOS
    return _ILEGAL.sub("", texto)[:MAX_CELDA], None


def _ancho(columna, titulo):
    c = columna.lower()
    if c in COLUMNAS_URL or c.startswith(("objeto", "descrip", "nombre_del_proc")):
        return 50
    if "valor" in c or "precio" in c:
        return 18
    if "fecha" in c:
        return 14
    return min(max(len(str(titulo)) + 2, 12), 40)


def _col_url(columnas):
    return next((c for c in COLUMNAS_URL if c in columnas), None)


class _Spool:
    """Archivo temporal JSONL: guarda filas en disco y recuerda la union de columnas."""

    def __init__(self):
        self._f = tempfile.TemporaryFile("w+", encoding="utf-8", newline="")
        self.columnas = {}
        self.n = 0

    def agregar(self, filas):
        for fila in filas:
            for k in fila:
                self.columnas.setdefault(k, None)
            self._f.write(json.dumps(fila, ensure_ascii=False, default=str) + "\n")
            self.n += 1

    def filas(self):
        self._f.seek(0)
        for linea in self._f:
            yield json.loads(linea)

    def cerrar(self):
        self._f.close()


def _revisar(cancelado, i):
    if i % LOTE_CANCELACION == 0 and cancelado():
        raise ExportacionCancelada()


def _escribir_csv(ruta, columnas, encabezados, filas, delimitador, cancelado):
    with open(ruta, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=delimitador, lineterminator="\n")
        w.writerow(encabezados)
        for i, fila in enumerate(filas):
            _revisar(cancelado, i)
            w.writerow([valor_crudo(fila.get(c)) for c in columnas])


def _escribir_json(ruta, columnas, filas, cancelado):
    with open(ruta, "w", encoding="utf-8") as f:
        f.write("[")
        for i, fila in enumerate(filas):
            _revisar(cancelado, i)
            registro = {c: valor_crudo(fila.get(c)) for c in columnas}
            f.write(("," if i else "") + "\n  " + json.dumps(registro, ensure_ascii=False))
        f.write("\n]\n")


def _hoja_xlsx(wb, nombre, columnas, encabezados, filas, col_url, cancelado):
    ws = wb.create_sheet(nombre)
    for i, (col, titulo) in enumerate(zip(columnas, encabezados), start=1):
        ws.column_dimensions[get_column_letter(i)].width = _ancho(col, titulo)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columnas))}1048576"
    cabecera = []
    for titulo in encabezados:
        celda = WriteOnlyCell(ws, value=_ILEGAL.sub("", str(titulo)))
        celda.font = NEGRITA
        cabecera.append(celda)
    ws.append(cabecera)
    for i, fila in enumerate(filas):
        _revisar(cancelado, i)
        celdas = []
        for col in columnas:
            valor, fmt = convertir_celda(col, fila.get(col))
            celda = WriteOnlyCell(ws, value=valor)
            if fmt:
                celda.number_format = fmt
            if (col == col_url and isinstance(valor, str)
                    and valor.startswith(("http://", "https://"))):
                celda.hyperlink = valor
                celda.font = ENLACE
            celdas.append(celda)
        ws.append(celdas)


def _cerrar_libro(wb):
    """Cierra los generadores de filas pendientes de openpyxl mientras su archivo
    sigue abierto (evita 'Exception ignored' al recolectarlos tras wb.close())."""
    for ws in wb._sheets:
        filas = getattr(ws, "_rows", None)
        if filas is not None and hasattr(filas, "close"):
            try:
                filas.close()
            except Exception:
                pass
    wb.close()


def _ruta_csv(carpeta, nombre, sello, creados):
    limpio = re.sub(r"[^\w\-]+", "_", nombre, flags=re.UNICODE).strip("_")[:40] or "datos"
    ruta, i = os.path.join(carpeta, f"SECOP_{limpio}_{sello}.csv"), 2
    while ruta in creados or os.path.exists(ruta):
        ruta = os.path.join(carpeta, f"SECOP_{limpio}_{sello}_{i}.csv")
        i += 1
    return ruta


def exportar_incremental(formato, columnas, destino, paginas, titulos=None,
                         delimitador=",", sello="", cancelado=None, hoja="Resumen"):
    """Escribe `paginas` (iterable de (filas, detalle)) sin materializarlas.
    `hoja`: nombre de la hoja (xlsx) o del archivo (csv_dataset) de las filas principales.
    Devuelve (rutas creadas, n_filas_del_resumen)."""
    if not columnas:
        raise ValueError("Seleccione al menos una columna.")
    if formato not in FORMATOS:
        raise ValueError(f"Formato desconocido: {formato}")
    cancelado = cancelado or (lambda: False)
    titulos = titulos or {}
    resumen, detalles, partes, finales, wb = _Spool(), {}, [], [], None
    try:
        for filas, detalle in paginas:
            if cancelado():
                raise ExportacionCancelada()
            resumen.agregar(filas)
            for nombre, datos in (detalle or {}).items():
                if datos:
                    if nombre not in detalles:
                        detalles[nombre] = _Spool()
                    detalles[nombre].agregar(datos)
        if cancelado():
            raise ExportacionCancelada()
        if resumen.n == 0:
            raise ValueError("No hay filas para exportar.")
        enc = [titulos.get(c, c) for c in columnas]

        if formato == "xlsx":
            parte = destino + ".part"
            partes.append(parte)
            wb = Workbook(write_only=True)
            usados = set()
            _hoja_xlsx(wb, nombre_hoja(hoja, usados), columnas, enc, resumen.filas(),
                       _col_url(columnas), cancelado)
            for nombre, sp in detalles.items():
                cols = list(sp.columnas)
                _hoja_xlsx(wb, nombre_hoja(nombre, usados), cols, cols, sp.filas(),
                           _col_url(cols), cancelado)
            wb.save(parte)
            os.replace(parte, destino)
            finales.append(destino)
        elif formato in ("csv", "json"):
            parte = destino + ".part"
            partes.append(parte)
            if formato == "csv":
                _escribir_csv(parte, columnas, enc, resumen.filas(), delimitador, cancelado)
            else:
                _escribir_json(parte, columnas, resumen.filas(), cancelado)
            os.replace(parte, destino)
            finales.append(destino)
        else:  # csv_dataset: `destino` es una carpeta
            trabajos = [(hoja, columnas, enc, resumen)]
            trabajos += [(n, list(sp.columnas), list(sp.columnas), sp)
                         for n, sp in detalles.items()]
            for nombre, cols, encs, sp in trabajos:
                ruta = _ruta_csv(destino, nombre, sello, finales)
                parte = ruta + ".part"
                partes.append(parte)
                _escribir_csv(parte, cols, encs, sp.filas(), delimitador, cancelado)
                os.replace(parte, ruta)
                finales.append(ruta)
        return list(finales), resumen.n
    except BaseException:
        for ruta in partes + finales:
            try:
                os.remove(ruta)
            except OSError:
                pass
        raise
    finally:
        if wb is not None:
            _cerrar_libro(wb)
        cerrar = getattr(paginas, "close", None)
        if cerrar:
            cerrar()
        resumen.cerrar()
        for sp in detalles.values():
            sp.cerrar()


def exportar(formato, columnas, filas, destino, titulos=None, detalle=None,
             delimitador=",", sello="", hoja="Resumen"):
    """Atajo para datos ya en memoria (una sola pagina)."""
    archivos, _ = exportar_incremental(formato, columnas, destino,
                                       iter([(filas, detalle or {})]), titulos,
                                       delimitador, sello, hoja=hoja)
    return archivos
